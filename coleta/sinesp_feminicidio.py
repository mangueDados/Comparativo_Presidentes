"""Feminicídios no Brasil a partir das bases Sinesp VDE (Ministério da Justiça
e Segurança Pública), registros das polícias estaduais.

Uso, a partir da raiz do projeto (lê ~290 MB de planilhas; leva ~10 minutos):
    python coleta/sinesp_feminicidio.py

Brutos em dados/brutos/sinesp_vde/ (fora do git; ver MANIFESTO.csv, com URL,
data de acesso e sha256 de cada arquivo). Soma 'total_vitima' das linhas com
evento 'Feminicídio' (esfera estadual, todos os municípios, inclusive 'NÃO
INFORMADO'). Grava o total mensal em dados/derivados/ e os anos completos em
dados/indicadores.csv.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

try:
    from coleta.comum import INDICATORS_FILE, ROOT, merge_into
except ModuleNotFoundError:  # execução direta: python coleta/sinesp_feminicidio.py
    from comum import INDICATORS_FILE, ROOT, merge_into

RAW_DIR = ROOT / "dados" / "brutos" / "sinesp_vde"
MANIFEST = RAW_DIR / "MANIFESTO.csv"
MONTHLY_FILE = ROOT / "dados" / "derivados" / "sinesp_feminicidio_mensal.csv"
SOURCE_URL = (
    "https://www.gov.br/mj/pt-br/assuntos/sua-seguranca/seguranca-publica/estatistica/"
    "download/dnsp-base-de-dados/bancovde-{year}.xlsx/@@download/file"
)
ACCESS_DATE = "2026-10-07"
EVENT = "Feminicídio"
# Ano em consolidação: estados ainda enviam dados atrasados (ver ADR-015).
PRELIMINARY_YEARS = {2025}

INDICATOR = "Feminicídios (registros policiais)"
SERIES = "Brasil"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_manifest() -> None:
    with MANIFEST.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(("arquivo", "url", "data_acesso", "bytes", "sha256"))
        for path in sorted(RAW_DIR.glob("bancovde-*.xlsx")):
            year = path.name.split("-")[1][:4]
            writer.writerow((path.name, SOURCE_URL.format(year=year), ACCESS_DATE,
                             path.stat().st_size, sha256(path)))


def monthly_victims(path: Path) -> dict[str, dict[str, int]]:
    """{AAAA-MM: {'vitimas': n, 'feminino': n}} das linhas de feminicídio."""
    workbook = load_workbook(path, read_only=True)
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = {name: index for index, name in enumerate(next(rows))}
    totals: dict[str, dict[str, int]] = defaultdict(lambda: {"vitimas": 0, "feminino": 0})
    for row in rows:
        if row[header["evento"]] != EVENT:
            continue
        if row[header["abrangencia"]] != "Estadual":
            raise ValueError(f"{path.name}: feminicídio fora da esfera estadual.")
        month = row[header["data_referencia"]].strftime("%Y-%m")
        totals[month]["vitimas"] += row[header["total_vitima"]] or 0
        totals[month]["feminino"] += row[header["feminino"]] or 0
    workbook.close()
    return dict(totals)


def annual(monthly: dict[str, dict[str, int]]) -> dict[int, tuple[int, int, int]]:
    """{ano: (vítimas, das quais sexo feminino, meses com dado)}."""
    years: dict[int, list[int]] = defaultdict(lambda: [0, 0, 0])
    for month, values in monthly.items():
        year = years[int(month[:4])]
        year[0] += values["vitimas"]
        year[1] += values["feminino"]
        year[2] += 1
    return {year: tuple(values) for year, values in years.items()}


def to_rows(years: dict[int, tuple[int, int, int]]) -> list[dict[str, str]]:
    rows = []
    for year, (victims, female, months) in sorted(years.items()):
        if months != 12:
            continue  # ano incompleto não é comparável a anos completos
        rows.append({
            "indicador": INDICATOR,
            "serie": SERIES,
            "ano": str(year),
            "valor": str(victims),
            "unidade": "vítimas",
            "fonte": "MJSP, Sinesp VDE (registros das polícias estaduais)",
            "url_fonte": SOURCE_URL.format(year=year),
            "periodo_ref": f"jan–dez/{year}",
            "inicio_periodo": f"{year}-01-01",
            "fim_periodo": f"{year}-12-31",
            "tipo_dado": "preliminar" if year in PRELIMINARY_YEARS else "observado",
            "data_acesso": ACCESS_DATE,
            "validada": "nao",
            "metodologia": (
                f"Soma de 'total_vitima' do evento 'Feminicídio' em todos os municípios e meses; "
                f"{female} vítimas registradas como sexo feminino e {victims - female} sem essa "
                "marcação. A base é atualizada com envios atrasados dos estados, e os valores "
                "recentes mudam. Transformação em coleta/sinesp_feminicidio.py."
            ),
        })
    return rows


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--output", type=Path, default=INDICATORS_FILE)
    args = parser.parse_args()

    write_manifest()
    monthly: dict[str, dict[str, int]] = {}
    for path in sorted(RAW_DIR.glob("bancovde-*.xlsx")):
        monthly.update(monthly_victims(path))
        print(f"lido {path.name}", flush=True)

    MONTHLY_FILE.parent.mkdir(parents=True, exist_ok=True)
    with MONTHLY_FILE.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(("mes", "vitimas", "vitimas_sexo_feminino"))
        for month in sorted(monthly):
            writer.writerow((month, monthly[month]["vitimas"], monthly[month]["feminino"]))

    years = annual(monthly)
    rows = to_rows(years)
    merge_into(args.output, INDICATOR, SERIES, rows)
    incomplete = [f"{y} ({m} meses)" for y, (_, _, m) in sorted(years.items()) if m != 12]
    print(f"{len(rows)} anos gravados; excluídos por incompletos: {', '.join(incomplete) or 'nenhum'}.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
