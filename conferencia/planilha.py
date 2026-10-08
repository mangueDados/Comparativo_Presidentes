"""Gera a planilha de conferência no padrão Mangue a partir de
conferencia/resultado_conferencia.csv (criado por conferencia/verificar.py).

Uso, a partir da raiz do projeto:
    python conferencia/planilha.py
"""

from __future__ import annotations

import csv
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from conferencia.verificar import (  # noqa: E402
    LOG_FILE, MANUAL, OK, DIVERGE, RESULT_FILE, SERIES_PENDING,
)

LOGO = ROOT / "logo-300x75.png"
VINHO, CREME, SUB, TEXTO, AUSENTE, BORDA, ABA2 = (
    "420001", "FAF7F1", "5C3A3A", "333333", "9A8F8A", "E3D6CC", "D9BFB8"
)
DECISIONS = ("aprovar", "rejeitar")
REVIEW_COLUMNS = ("decisao", "conferido_por", "observacao")


def aplicar_padrao_mangue(ws, titulo, subtitulo="", principal=False, larguras=None,
                          cols_texto_longo=()):
    """Padrão Mangue para uma aba com cabeçalho na linha 5 (skill padrao-mangue-planilhas)."""
    ncol, nrow = ws.max_column, ws.max_row
    creme = PatternFill("solid", fgColor=CREME)
    fino = Side(style="thin", color=BORDA)
    for c in range(1, ncol + 1):
        for r in (1, 2, 3):
            ws.cell(r, c).fill = creme
        ws.cell(3, c).border = Border(bottom=Side(style="medium", color=VINHO))
    ws.row_dimensions[1].height = 46
    ws.row_dimensions[2].height = 26
    # O logo enviado tem 300×75; mantém a proporção na altura do padrão (77 px).
    img = XLImage(str(LOGO))
    img.width, img.height = 308, 77
    ws.add_image(img, "A1")
    ws["A2"] = titulo
    ws["A2"].font = Font(name="Arial", size=15, bold=True, color=VINHO)
    ws["A3"] = subtitulo
    ws["A3"].font = Font(name="Arial", size=10, italic=True, color=SUB)
    for c in range(1, ncol + 1):
        h = ws.cell(5, c)
        h.font = Font(name="Arial", bold=True, color=CREME)
        h.fill = PatternFill("solid", fgColor=VINHO)
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        h.border = Border(bottom=fino)
    for r in range(6, nrow + 1):
        for c in range(1, ncol + 1):
            cel = ws.cell(r, c)
            if cel.font.color is None or cel.font.color.rgb not in (f"00{AUSENTE}", f"00{VINHO}"):
                cel.font = Font(name="Arial", size=10, color=TEXTO)
            cel.border = Border(bottom=fino)
            if r % 2 == 1 and r > 6:
                cel.fill = creme
            if get_column_letter(c) in cols_texto_longo:
                cel.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:{get_column_letter(ncol)}{nrow}"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = VINHO if principal else ABA2
    for col, w in (larguras or {}).items():
        ws.column_dimensions[col].width = w


def ausente(cel, motivo="sem dados na fonte"):
    cel.value = motivo
    cel.font = Font(name="Arial", size=10, italic=True, color=AUSENTE)


def destaque(cel):
    cel.font = Font(name="Arial", size=10, bold=True, color=VINHO)


def _number(text: str):
    try:
        return float(text)
    except ValueError:
        return text


def _write_table(ws, header, rows):
    for c, name in enumerate(header, start=1):
        ws.cell(5, c, name)
    for r, row in enumerate(rows, start=6):
        for c, value in enumerate(row, start=1):
            ws.cell(r, c, value)


def build(output: Path) -> Path:
    if not LOGO.is_file():
        raise SystemExit(f"Logo da Mangue não encontrado em {LOGO}; não recrie o logo.")
    with RESULT_FILE.open(encoding="utf-8") as f:
        results = list(csv.DictReader(f))
    with LOG_FILE.open(encoding="utf-8") as f:
        log = list(csv.DictReader(f))

    wb = Workbook()

    # 1. Conferência (leitura obrigatória)
    ws = wb.active
    ws.title = "Conferência"
    header = ("indicador", "serie", "ano", "periodo_ref", "valor_csv", "valor_checagem",
              "diferenca", "resultado", "fonte_checagem", "como_conferir", "pendencia",
              *REVIEW_COLUMNS)
    _write_table(ws, header, [
        (r["indicador"], r["serie"], int(r["ano"]), r["periodo_ref"], _number(r["valor_csv"]),
         _number(r["valor_checagem"]) if r["valor_checagem"] else None,
         _number(r["diferenca"]) if r["diferenca"] else None, r["resultado"],
         r["fonte_checagem"], r["como_conferir"], r["pendencia"] or None, None, None, None)
        for r in results
    ])
    counts = {k: sum(r["resultado"] == k for r in results) for k in (OK, DIVERGE, MANUAL)}
    aplicar_padrao_mangue(
        ws, "Conferência dos dados do comparativo",
        f"{counts[OK]} OK na checagem independente, {counts[DIVERGE]} divergentes, "
        f"{counts[MANUAL]} para conferir à mão. Preencha 'decisao' (aprovar/rejeitar) e "
        "'conferido_por'; depois rode conferencia/aplicar.py.",
        principal=True,
        larguras={"A": 34, "B": 14, "C": 8, "D": 28, "E": 12, "F": 14, "G": 11, "H": 12,
                  "I": 30, "J": 48, "K": 48, "L": 12, "M": 18, "N": 36},
        cols_texto_longo=("A", "D", "I", "J", "K", "N"),
    )
    for r in range(6, ws.max_row + 1):
        if ws.cell(r, 6).value is None:
            ausente(ws.cell(r, 6), "sem checagem automática")
        if ws.cell(r, 7).value is None:
            ausente(ws.cell(r, 7), "–")
        if ws.cell(r, 8).value != OK:
            destaque(ws.cell(r, 8))
        if ws.cell(r, 11).value:
            destaque(ws.cell(r, 11))
        else:
            ausente(ws.cell(r, 11), "nenhuma")
    choice = DataValidation(type="list", formula1=f'"{",".join(DECISIONS)}"', allow_blank=True)
    ws.add_data_validation(choice)
    choice.add(f"L6:L{ws.max_row}")

    # 2. Pendências (leitura obrigatória)
    ws = wb.create_sheet("Pendências")
    _write_table(ws, ("serie", "o_que_decidir"), sorted(SERIES_PENDING.items()))
    aplicar_padrao_mangue(
        ws, "Pendências da série inteira",
        "Decisões editoriais que não dependem de um valor específico; resolver antes de publicar.",
        principal=True, larguras={"A": 40, "B": 100}, cols_texto_longo=("A", "B"),
    )

    # 3. Metodologia
    ws = wb.create_sheet("Metodologia")
    _write_table(ws, ("item", "explicação"), [
        ("Situação", f"Gerada em {date.today():%d/%m/%Y} a partir de dados/indicadores.csv; "
                     "nenhuma linha está aprovada até a decisão humana nesta planilha."),
        ("Universo", f"{len(results)} valores de 8 séries: PRODES Amazônia, IPCA, PIB, PIB per "
                     "capita, taxa de homicídios, mortes por intervenção legal, parcela delas nos "
                     "homicídios e mortes de causa indeterminada."),
        ("Checagem independente", "Cada valor foi baixado de novo por outro caminho: IBGE pela API "
                                  "SIDRA 'values'; Atlas da Violência pelo SIM/DATASUS (TabNet), "
                                  "sua fonte primária; PRODES pelo texto de notícias do INPE."),
        ("OK", "Diferença absoluta até 0,006 em relação à checagem (arredondamento da fonte)."),
        ("DIVERGE", "Diferença maior que a tolerância: não aprovar sem explicar em 'observacao'."),
        ("MANUAL", "Sem segunda fonte automática; conferir pelo link de 'como_conferir'."),
        ("Homicídios no TabNet", "Atlas = agressões (X85-Y09) + intervenções legais (Y35-Y36), "
                                 "por residência; a soma bateu com o Atlas em todos os anos."),
        ("Taxa de homicídios", "Recalculada como óbitos ÷ população do IBGE (tabela 6784) × "
                               "100 mil; 2024 sem população na tabela, por isso MANUAL."),
        ("PRODES", "A base geográfica do INPE mede incremento mapeado, não a taxa oficial; "
                   "por isso não foi usada como checagem. Anos sem texto oficial ficam MANUAL."),
        ("Sugestão de conferência", "Conferir à mão todas as linhas MANUAL e com pendência e, nas "
                                    "linhas OK, uma amostra (por exemplo, o primeiro e o último ano "
                                    "de cada série)."),
        ("Aplicação", "conferencia/aplicar.py marca validada=sim só nas linhas 'aprovar' com nome "
                      "em 'conferido_por' e cujo valor não mudou desde a geração da planilha."),
    ])
    aplicar_padrao_mangue(ws, "Metodologia", "Como a conferência foi feita e como ler o resultado.",
                          larguras={"A": 26, "B": 100}, cols_texto_longo=("B",))

    # 4. Log de URLs
    ws = wb.create_sheet("Log de URLs")
    _write_table(ws, ("url", "params", "status", "quando"),
                 [(r["url"], r["params"], r["status"], r["quando"]) for r in log])
    aplicar_padrao_mangue(ws, "Log de URLs", "Requisições feitas pela checagem automática.",
                          larguras={"A": 90, "B": 10, "C": 12, "D": 18}, cols_texto_longo=("A",))

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output


def main() -> int:
    output = ROOT / "conferencia" / f"conferencia_comparativo_v1_{date.today():%Y%m%d}.xlsx"
    print(f"Planilha: {build(output).relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
